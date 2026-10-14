#!/usr/bin/env python3
"""Lấy số kỳ chưa về Tam hoa Bingo18 và lưu lịch sử vào Excel trên GitHub."""

from __future__ import annotations

import argparse
import json
import os
import ssl
import sys
import time
import urllib.error
import urllib.request
from datetime import date, datetime, time as datetime_time, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


DATA_URL = "https://www.bingo18.top/data/data.json"
WORKBOOK_NAME = "Bingo18_Lich_su_Tam_hoa.xlsx"
SHEET_NAME = "Lịch sử"
TRIPLES = ["111", "222", "333", "444", "555", "666"]
VN_TIMEZONE = ZoneInfo("Asia/Ho_Chi_Minh")



SCHEDULE_TIME = datetime_time(22, 5)


def resolve_statistics_date(
    now: datetime, event_name: str, explicit_date: str | None = None
) -> date:
    """Scheduled jobs use the latest 22:05 session; manual jobs use today.

    --statistics-date or BINGO18_STATISTICS_DATE (YYYY-MM-DD) overrides
    this rule for a manual backfill or a rerun of an older session.
    The cutoff fallback cannot identify delays longer than one daily cycle.
    """
    now = now.astimezone(VN_TIMEZONE)
    if explicit_date is not None:
        return date.fromisoformat(explicit_date)
    if event_name == "schedule" and now.time() < SCHEDULE_TIME:
        return now.date() - timedelta(days=1)
    return now.date()


def fetch_json() -> tuple[dict, str]:
    request = urllib.request.Request(
        DATA_URL,
        headers={
            "User-Agent": "Mozilla/5.0 Bingo18History/2.0",
            "Accept": "application/json,text/plain,*/*",
        },
    )
    last_error = None
    for attempt in range(1, 4):
        try:
            with urllib.request.urlopen(request, timeout=45) as response:
                return json.load(response), "Kết nối bảo mật bình thường"
        except urllib.error.URLError as exc:
            last_error = exc
            if "CERTIFICATE_VERIFY_FAILED" in str(exc):
                with urllib.request.urlopen(
                    request, timeout=45, context=ssl._create_unverified_context()
                ) as response:
                    return json.load(response), "Chế độ tương thích chứng chỉ"
            if attempt < 3:
                time.sleep(5)
    raise last_error


def extract_counts(data: dict) -> dict[str, int]:
    stats = data.get("gbingoStatistics")
    if not isinstance(stats, list):
        raise ValueError("Dữ liệu không có mục gbingoStatistics.")
    by_type = {item.get("typePlay"): item for item in stats if isinstance(item, dict)}
    result: dict[str, int] = {}
    for digit in range(1, 7):
        item = by_type.get(f"TRIPLE_{digit}")
        value = item.get("countNearest") if item else None
        if not isinstance(value, int) or value < 0:
            raise ValueError(f"Không có dữ liệu hợp lệ cho {digit}{digit}{digit}.")
        result[str(digit) * 3] = value
    return result


def write_history(
    path: Path, counts: dict[str, int], note: str,
    statistics_date: date | None = None,
) -> str:
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy file Excel mẫu: {path}")
    wb = load_workbook(path)
    ws = wb[SHEET_NAME]
    now = datetime.now(VN_TIMEZONE)
    if statistics_date is None:
        statistics_date = resolve_statistics_date(
            now, os.environ.get("GITHUB_EVENT_NAME", ""),
            os.environ.get("BINGO18_STATISTICS_DATE"),
        )

    target_row = None
    for row in range(5, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, datetime):
            value = value.date()
        if value == statistics_date:
            target_row = row
            break

    action = "Cập nhật dòng ngày thống kê" if target_row else "Thêm dòng mới"
    if target_row is None and ws.cell(5, 1).value is None:
        target_row = 5
    target_row = target_row or max(ws.max_row + 1, 5)

    values = [
        datetime.combine(statistics_date, datetime_time.min),
        now.replace(tzinfo=None),
        *[counts[key] for key in TRIPLES],
        DATA_URL,
        f"Thành công - {note}",
    ]
    for col, value in enumerate(values, start=1):
        ws.cell(target_row, col, value)
    ws.cell(target_row, 1).number_format = "dd/mm/yyyy"
    ws.cell(target_row, 2).number_format = "dd/mm/yyyy hh:mm:ss"
    for col in range(3, 9):
        ws.cell(target_row, col).number_format = "#,##0"
        ws.cell(target_row, col).alignment = Alignment(horizontal="center")

    max_value = max(counts.values())
    for col, key in enumerate(TRIPLES, start=3):
        cell = ws.cell(target_row, col)
        is_max = counts[key] == max_value
        cell.fill = PatternFill("solid", fgColor="FDE9D9" if is_max else "FFFFFF")
        cell.font = Font(bold=is_max, color="C00000" if is_max else "1F2937")

    wb.save(path)
    return action


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--statistics-date", default=os.environ.get("BINGO18_STATISTICS_DATE"),
        help="Ngày thống kê chạy bù, YYYY-MM-DD (ưu tiên hơn lịch tự động).",
    )
    args = parser.parse_args()
    try:
        statistics_date = resolve_statistics_date(
            datetime.now(VN_TIMEZONE), os.environ.get("GITHUB_EVENT_NAME", ""),
            args.statistics_date,
        )
    except ValueError:
        parser.error("Ngày thống kê phải là ngày hợp lệ theo định dạng YYYY-MM-DD.")
    path = Path(__file__).resolve().parent / WORKBOOK_NAME
    data, note = fetch_json()
    counts = extract_counts(data)
    action = write_history(path, counts, note, statistics_date)
    print(action + f" ({statistics_date:%d/%m/%Y}): " + ", ".join(f"{k}={counts[k]}" for k in TRIPLES))


if __name__ == "__main__":
    main()
